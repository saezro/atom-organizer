"""Consenso de giro en el manifiesto: columna `giro_voto`, tabla `vuelos`, migración y Excel."""
import sqlite3
import types

from openpyxl import load_workbook

from atom_core import indice as indice_mod
from atom_core import indice_excel
from atom_core.manifiesto import FilaManifiesto, Manifiesto

V = indice_mod.VERSION_ALGORITMO_GIRO


def _vuelo(**kw):
    base = dict(pb="1", vuelo="2", sufijo="", angulo=90, n_total=10, n_90=9, n_270=1, n_0=0,
                pct_90=90.0, pct_270=10.0, pct_0=0.0, pct_ganadora=90.0, estado="ok",
                motivo=None, version_algoritmo=V)
    base.update(kw)
    return base


def _fila(nombre="a.JPG", **kw):
    base = dict(ruta_origen=f"/x/{nombre}", tipo="RGB", timestamp_exif="2026-01-01T10:00:00",
                pb="1", vuelo="2", nombre_nuevo=nombre, angulo_giro=90, angulo_version=V,
                ruta_salida_original=f"/o/{nombre}", giro_voto=90, modelo=None,
                pct_recorte=None, comprime=False, ruta_salida_crop=None, ruta_salida_tiff=None,
                unassigned=False)
    base.update(kw)
    return FilaManifiesto(**base)


def test_migracion_manifiesto_viejo_idempotente(tmp_path):
    db = tmp_path / "m.db"
    con = sqlite3.connect(db)
    con.execute("""CREATE TABLE imagenes (id INTEGER PRIMARY KEY AUTOINCREMENT, ruta_origen TEXT NOT NULL,
        nombre_original TEXT NOT NULL DEFAULT '', clave TEXT NOT NULL, tipo TEXT NOT NULL,
        timestamp_exif TEXT, modelo TEXT, make TEXT, pb TEXT, vuelo TEXT, equipo_estadillo TEXT,
        nombre_nuevo TEXT NOT NULL DEFAULT '', angulo_giro INTEGER NOT NULL DEFAULT 0, pct_recorte REAL,
        comprime INTEGER NOT NULL DEFAULT 0, ruta_salida_original TEXT NOT NULL, ruta_salida_crop TEXT,
        ruta_salida_tiff TEXT, unassigned INTEGER NOT NULL DEFAULT 0, estado TEXT NOT NULL DEFAULT 'pendiente',
        motivo_fallo TEXT, verificacion TEXT, ejecucion_id INTEGER)""")
    con.execute("INSERT INTO imagenes (ruta_origen, clave, tipo, pb, vuelo, angulo_giro, ruta_salida_original)"
                " VALUES ('/x/a.JPG','k','RGB','1','2',90,'/o/a.JPG')")
    con.commit()
    con.close()
    for _ in range(2):
        m = Manifiesto(db)
        m.crear_esquema()
        cols = [r["name"] for r in m._conexion().execute("PRAGMA table_info(imagenes)")]
        assert "giro_voto" in cols
        assert m.vuelos() == {}
        # sin fila en vuelos: cae a imagenes (versión None = antigua)
        assert m.angulos_versionados_por_vuelo() == {("1", "2"): (90, None)}
        m.cerrar()


def test_insertar_con_vuelos_y_lectura(tmp_path):
    m = Manifiesto(tmp_path / "m.db")
    m.crear_esquema()
    m.insertar_o_reabrir([_fila()], vuelos=[_vuelo()])
    assert m.todas()[0]["giro_voto"] == 90
    assert m.vuelos()[("1", "2", None)]["estado"] == "ok"
    assert m.angulos_versionados_por_vuelo()[("1", "2")] == (90, V)
    m.guardar_vuelos([_vuelo(angulo=0, estado="sin_consenso", version_algoritmo="vieja")])
    assert m.angulos_versionados_por_vuelo()[("1", "2")] == (0, "vieja")
    m.cerrar()


def test_hoja_vuelos_en_excel(tmp_path):
    destino = tmp_path / "PLANTA"
    (destino / ".organizado").mkdir(parents=True)
    m = Manifiesto(destino / ".organizado" / "m.db")
    m.crear_esquema()
    m.insertar_o_reabrir([_fila()], vuelos=[_vuelo(estado="sin_consenso", motivo="x")])
    cfg = types.SimpleNamespace(output_folder=str(destino), flight_height=50.0,
                                calculate_proyected_distance=False)
    ruta = indice_excel.escribir_indice(m, cfg, {}, types.SimpleNamespace(emit=lambda *a, **k: None))
    libro = load_workbook(ruta)
    assert libro.sheetnames == ["Imagenes", "Vuelos"]
    cab = [c.value for c in libro["Vuelos"][1]]
    fila = dict(zip(cab, [c.value for c in libro["Vuelos"][2]]))
    assert fila["Estado"] == "sin_consenso" and fila["N90"] == 9
    cab_i = [c.value for c in libro["Imagenes"][1]]
    assert libro["Imagenes"][2][cab_i.index("GiroVoto")].value == 90
    m.cerrar()
