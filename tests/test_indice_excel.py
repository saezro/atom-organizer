import os
from types import SimpleNamespace

from openpyxl import load_workbook

from atom_core import indice_excel
from atom_core.manifiesto import FilaManifiesto, Manifiesto


class _Cb:
    def __init__(self):
        self.lineas = []

    def emit(self, v):
        self.lineas.append(v)


def _preparar(tmp_path):
    destino = tmp_path / "MELINESTI"
    (destino / ".organizado").mkdir(parents=True)
    m = Manifiesto(destino / ".organizado" / "manifiesto.db")
    m.crear_esquema()
    ej = m.abrir_ejecucion("/media/sd1", "3.4.93")
    orig = str(destino / "RGB" / "PB1" / "PB1_V1" / "20260115_100000_DJI_0001_D.JPG")
    m.insertar_o_reabrir([
        FilaManifiesto(ruta_origen="/media/sd1/DJI_0001_D.JPG", tipo="RGB",
                       timestamp_exif="2026-01-15T10:00:00", modelo="M4T", make="DJI",
                       equipo_estadillo="DJI M300", pb="1", vuelo="1",
                       nombre_nuevo="20260115_100000_DJI_0001_D.JPG", angulo_giro=270,
                       pct_recorte=0.8, comprime=True, ruta_salida_original=orig,
                       ruta_salida_crop=orig.replace(".JPG", "_CROP.JPG"), ruta_salida_tiff=None,
                       unassigned=False, bytes_origen=9000, meta_leida=True, lat=37.1, lon=-5.6,
                       gimbal_yaw="+12.50", gimbal_pitch="-90.00", altura_relativa="+50.000",
                       ancho_px=4032, alto_px=3024),
        FilaManifiesto(ruta_origen="/media/sd1/DJI_9999_D.JPG", tipo="RGB", timestamp_exif=None,
                       modelo=None, pb=None, vuelo=None, nombre_nuevo="", angulo_giro=0,
                       pct_recorte=None, comprime=True,
                       ruta_salida_original=str(destino / "SIN_ORDENAR" / "RGB" / "DJI_9999_D.JPG"),
                       ruta_salida_crop=None, ruta_salida_tiff=None, unassigned=True),
    ], ejecucion_id=ej)
    filas = m.todas()
    m.marcar_hecha(filas[0]["id"], "")
    m.marcar_fallida(filas[1]["id"], "sin timestamp")
    cfg = SimpleNamespace(output_folder=str(destino), flight_height=50.0,
                          calculate_proyected_distance=True)
    return m, cfg, orig


def test_escribe_una_fila_por_imagen_con_columnas_comunes(tmp_path):
    m, cfg, orig = _preparar(tmp_path)
    ruta = indice_excel.escribir_indice(m, cfg, {orig: (0.0, 37.1, -5.6)}, _Cb())

    assert ruta == os.path.join(cfg.output_folder, "INDICE_MELINESTI.xlsx")
    ws = load_workbook(ruta, read_only=True)["Imagenes"]
    filas = list(ws.iter_rows(values_only=True))
    assert filas[0] == indice_excel.COLUMNAS
    assert len(filas) == 3
    fila = dict(zip(indice_excel.COLUMNAS, filas[1]))
    assert fila["Vuelo"] == "1" and fila["Tipo"] == "RGB" and fila["SinOrdenar"] == "No"
    assert fila["NombreOriginal"] == "DJI_0001_D.JPG"
    assert fila["TimestampEXIF"] == "2026-01-15 10:00:00"
    assert (fila["Make"], fila["Model"], fila["EquipoEstadillo"], fila["EquipoCoincide"]) == ("DJI", "M4T", "DJI M300", "No")
    assert (fila["Lat"], fila["GimbalYaw"], fila["AlturaRelativa"]) == (37.1, 12.5, 50.0)
    assert (fila["AnguloGiro"], fila["PctRecorte"], fila["Comprime"]) == (270, 80.0, "Sí")
    assert fila["RutaOriginal"] == os.path.join("RGB", "PB1", "PB1_V1", "20260115_100000_DJI_0001_D.JPG")
    assert fila["CalculatedDistance"] == 0.0 and fila["AlturaVuelo"] == 50.0
    assert fila["Ejecucion"] == 1 and fila["FechaEjecucion"]
    fallida = dict(zip(indice_excel.COLUMNAS, filas[2]))
    assert (fallida["Estado"], fallida["MotivoFallo"], fallida["SinOrdenar"]) == ("fallido", "sin timestamp", "Sí")
    assert (fallida["Make"], fallida["EquipoCoincide"]) == (None, None)


def test_excel_abierto_escribe_con_otro_nombre_y_avisa(tmp_path, monkeypatch):
    m, cfg, _orig = _preparar(tmp_path)
    bloqueado = indice_excel.ruta_indice(cfg.output_folder)
    replace_real = os.replace

    def _replace(origen, destino):
        if os.path.abspath(destino) == os.path.abspath(bloqueado):
            raise PermissionError("abierto en Excel")
        return replace_real(origen, destino)

    monkeypatch.setattr(indice_excel.os, "replace", _replace)
    cb = _Cb()
    ruta = indice_excel.escribir_indice(m, cfg, {}, cb)

    assert ruta != bloqueado
    assert os.path.basename(ruta).startswith("INDICE_MELINESTI_") and ruta.endswith(".xlsx")
    assert os.path.isfile(ruta)
    assert any("abierto" in l for l in cb.lineas)
    assert not [f for f in os.listdir(cfg.output_folder) if f.endswith(".tmp")]
