"""Índice Excel de la planta: `<destino>/INDICE_<PLANTA>.xlsx`, una fila por
imagen del manifiesto.

Se regenera ENTERO en cada cierre desde el manifiesto (la fuente); nunca se
lee de vuelta. Organizar por cachitos sobre el mismo destino hace que el
Excel acumule solo, porque el manifiesto acumula."""
from __future__ import annotations

import datetime
import os
import tempfile

from atom_core.almacen import es_uri_gcs, publicar_en
from atom_core.equipo import equipo_coincide, texto_coincide

COLUMNAS = (
    "PB", "Vuelo", "Tipo", "SinOrdenar",
    "NombreOriginal", "NombreNuevo", "TimestampEXIF", "Make", "Model", "EquipoEstadillo", "EquipoCoincide",
    "AnchoPx", "AltoPx", "BytesOrigen",
    "Lat", "Lon", "AltitudAbs", "AlturaRelativa", "GimbalYaw", "GimbalPitch", "GimbalRoll",
    "FlightYaw", "AlturaVuelo", "CalculatedDistance", "LatitudFoto", "LongitudFoto",
    "AnguloGiro", "GiroVoto", "PctRecorte", "Comprime", "RutaOriginal", "RutaCrop", "RutaTIFF",
    "Estado", "MotivoFallo", "Ejecucion", "FechaEjecucion", "RutaOrigen",
)


COLUMNAS_VUELOS = (
    "PB", "Vuelo", "Sufijo", "Angulo", "NTotal", "N90", "N270", "N0",
    "Pct90", "Pct270", "Pct0", "PctGanadora", "Estado", "Motivo", "VersionAlgoritmo",
    "NT", "NTiff", "NW", "NCrop", "FilasMeta", "FilasLocation",
    "EstadoCoherencia", "MotivoCoherencia",
)
_N_GIRO = 15  # las primeras columnas de COLUMNAS_VUELOS que vienen del consenso de giro
_CLAVES_VUELOS = ("pb", "vuelo", "sufijo", "angulo", "n_total", "n_90", "n_270", "n_0",
                  "pct_90", "pct_270", "pct_0", "pct_ganadora", "estado", "motivo",
                  "version_algoritmo")


def nombre_planta(output_folder: str) -> str:
    return os.path.basename(str(output_folder).rstrip("/\\")) or "PLANTA"


def ruta_indice(output_folder: str) -> str:
    nombre = f"INDICE_{nombre_planta(output_folder)}.xlsx"
    if es_uri_gcs(output_folder):
        return f"{output_folder.rstrip('/')}/{nombre}"
    return os.path.join(output_folder, nombre)


def _numero(valor):
    """El XMP llega como texto (`+12.50`): en el Excel va como número para
    poder filtrar y ordenar. Lo que no sea número se deja tal cual."""
    if valor is None or valor == "":
        return None
    try:
        return float(valor)
    except (TypeError, ValueError):
        return valor


def _relativa(ruta: str | None, base: str) -> str | None:
    if not ruta:
        return None
    if es_uri_gcs(base):
        prefijo = base.rstrip("/") + "/"
        return ruta[len(prefijo):] if ruta.startswith(prefijo) else ruta
    try:
        return os.path.relpath(ruta, base)
    except ValueError:  # otra unidad en Windows
        return ruta


def _si_no(valor) -> str:
    return "Sí" if valor else "No"


def _filas(manifiesto, cfg, proyecciones):
    ejecuciones = manifiesto.ejecuciones()
    base = cfg.output_folder
    for fila in manifiesto.todas():
        ejecucion = ejecuciones.get(fila["ejecucion_id"])
        distancia, lat_foto, lon_foto = proyecciones.get(fila["ruta_salida_original"], (None, None, None))
        pct = fila["pct_recorte"]
        yield (
            fila["pb"], fila["vuelo"], fila["tipo"], _si_no(fila["unassigned"]),
            fila["nombre_original"], fila["nombre_nuevo"] or None,
            fila["timestamp_exif"].replace("T", " ") if fila["timestamp_exif"] else None,
            fila["make"], fila["modelo"], fila["equipo_estadillo"],
            texto_coincide(equipo_coincide(fila["equipo_estadillo"], fila["modelo"])),
            fila["ancho_px"], fila["alto_px"], fila["bytes_origen"],
            fila["lat"], fila["lon"], _numero(fila["altitud_abs"]), _numero(fila["altura_relativa"]),
            _numero(fila["gimbal_yaw"]), _numero(fila["gimbal_pitch"]), _numero(fila["gimbal_roll"]),
            _numero(fila["flight_yaw"]),
            cfg.flight_height if cfg.calculate_proyected_distance else None,
            _numero(distancia), _numero(lat_foto), _numero(lon_foto),
            fila["angulo_giro"], fila["giro_voto"], round(pct * 100, 2) if pct is not None else None, _si_no(fila["comprime"]),
            _relativa(fila["ruta_salida_original"], base), _relativa(fila["ruta_salida_crop"], base),
            _relativa(fila["ruta_salida_tiff"], base),
            fila["estado"], fila["motivo_fallo"], fila["ejecucion_id"],
            ejecucion["inicio"] if ejecucion is not None else None, fila["ruta_origen"],
        )


def filas_hoja_vuelos(vuelos_giro: dict, recuentos: dict) -> list[list]:
    """Una fila por vuelo del manifiesto (con o sin consenso de giro)."""
    giro_por_vuelo: dict = {}
    for (pb, vuelo, sufijo), v in vuelos_giro.items():
        giro_por_vuelo.setdefault((pb, vuelo), []).append((sufijo or "", v))
    filas = []
    for clave in sorted(set(giro_por_vuelo) | set(recuentos)):
        r = recuentos.get(clave, {})
        cola = [r.get("n_T"), r.get("n_tiff"), r.get("n_W"), r.get("n_crop"),
                r.get("filas_meta"), r.get("filas_location"),
                r.get("estado"), r.get("motivo") or None]
        for sufijo, v in sorted(giro_por_vuelo.get(clave, [("", None)]), key=lambda t: t[0]):
            if v is None:
                giro = [clave[0], clave[1]] + [None] * (_N_GIRO - 2)
            else:
                giro = [v[c] if c != "sufijo" else (v[c] or None) for c in _CLAVES_VUELOS]
            filas.append(giro + cola)
    return filas


def escribir_indice(manifiesto, cfg, proyecciones: dict, progress_callback, recuentos=None) -> str:
    """Escribe el índice y devuelve la ruta final. Si el fichero está abierto
    (Excel en Windows bloquea el `.xlsx`), escribe al lado con marca de hora y
    avisa: un índice bloqueado no puede tumbar el cierre."""
    from openpyxl import Workbook  # noqa: PLC0415 (import perezoso: falta en el .exe -> no debe tumbar el import de phases)
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    libro = Workbook(write_only=True)
    hoja = libro.create_sheet("Imagenes")
    hoja.freeze_panes = "A2"
    hoja.append(list(COLUMNAS))
    n_filas = 1
    for fila in _filas(manifiesto, cfg, proyecciones):
        hoja.append(list(fila))
        n_filas += 1

    if recuentos is None:
        from atom_core.cierre import calcular_recuentos  # noqa: PLC0415 (perezoso: cierre arrastra utils)
        recuentos = calcular_recuentos(manifiesto, cfg, None)
    from openpyxl.cell import WriteOnlyCell  # noqa: PLC0415
    from openpyxl.styles import Font, PatternFill  # noqa: PLC0415
    from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: PLC0415
    import warnings  # noqa: PLC0415

    hoja_vuelos = libro.create_sheet("Vuelos")
    hoja_vuelos.freeze_panes = "A2"

    def _celda(valor, fondo=None, color=None, negrita=False):
        c = WriteOnlyCell(hoja_vuelos, value=valor)
        if fondo:
            c.fill = PatternFill("solid", fgColor=fondo)
        if color or negrita:
            c.font = Font(color=color, bold=negrita)
        return c

    hoja_vuelos.append([_celda(n, "EE763C", "FFFFFF", True) for n in COLUMNAS_VUELOS])
    col_estado = COLUMNAS_VUELOS.index("EstadoCoherencia")
    n_vuelos = 1
    for fila in filas_hoja_vuelos(manifiesto.vuelos(), recuentos):
        celdas = []
        for i, v in enumerate(fila):
            if i == col_estado and v == "OK":
                celdas.append(_celda(v, "C6EFCE", "006100"))
            elif i == col_estado and v == "ERROR":
                celdas.append(_celda(v, "FFC7CE", "9C0006"))
            else:  # incluye EXCLUIDO: neutro, sin relleno
                celdas.append(_celda(v))
        hoja_vuelos.append(celdas)
        n_vuelos += 1
    tabla = Table(displayName="TablaVuelos",
                  ref=f"A1:{get_column_letter(len(COLUMNAS_VUELOS))}{max(n_vuelos, 2)}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    tabla._initialise_columns()
    for columna, nombre in zip(tabla.tableColumns, COLUMNAS_VUELOS):
        columna.name = nombre
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl avisa SIEMPRE en write-only aunque las columnas estén puestas
        hoja_vuelos.add_table(tabla)

    # Cabecera congelada + autofiltro (spec). `write_only=True` sí soporta
    # `auto_filter.ref` (openpyxl 3.1.5) mientras se fije antes de `save`.
    ultima_columna = get_column_letter(len(COLUMNAS))
    hoja.auto_filter.ref = f"A1:{ultima_columna}{n_filas}"

    destino = ruta_indice(cfg.output_folder)
    if es_uri_gcs(destino):
        descriptor, temporal = tempfile.mkstemp(suffix=".xlsx")
        os.close(descriptor)
        try:
            libro.save(temporal)
            publicar_en(temporal, destino)
        finally:
            os.unlink(temporal)
        return destino

    # Temporal en la MISMA carpeta: `os.replace` entre discos falla (EXDEV).
    temporal = destino + ".tmp"
    libro.save(temporal)
    try:
        os.replace(temporal, destino)
        return destino
    except PermissionError:
        marca = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        alternativo = os.path.join(os.path.dirname(destino),
                                   f"INDICE_{nombre_planta(cfg.output_folder)}_{marca}.xlsx")
        os.replace(temporal, alternativo)
        progress_callback.emit(
            f"\nAVISO: {os.path.basename(destino)} está abierto; el índice se ha "
            f"guardado como {os.path.basename(alternativo)}.\n")
        return alternativo
    finally:
        if os.path.exists(temporal):
            os.unlink(temporal)
